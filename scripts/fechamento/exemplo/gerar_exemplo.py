# -*- coding: utf-8 -*-
"""
Gera uma planilha de fechamento 100% INVENTADA, com o mesmo layout da aba
de mes do fechamento real (2025-2026), so para testar o enviar_fechamento.py
e a tela /fechamento.html sem usar nenhum dado real.

    python gerar_exemplo.py                 -> Fechamento_Exemplo.xlsx
    python gerar_exemplo.py --sem-valores   -> Fechamento_Exemplo_SemValores.xlsx
                                               (formulas sem valor calculado,
                                               pra testar o aviso do script)

Todos os nomes, valores, fornecedores e "documentos" daqui sao ficticios.
"""

import argparse
import datetime as dt
import os
import random

import openpyxl

PASTA = os.path.dirname(os.path.abspath(__file__))

MESES_NOME = ['Janeiro', 'Fevereiro', 'Marco', 'Abril', 'Maio', 'Junho', 'Julho', 'Agosto',
              'Setembro', 'Outubro', 'Novembro', 'Dezembro']

# (nome do produto, categoria, tipo de variacao, variacoes, subgrupos)
PRODUTOS = [
    ('Tapete pega areia', 'Pega areia', 'Cor', ['Amarelo', 'Azul', 'Preto', 'Rosa', 'Cinza'], ['70X50', '100X50', '120X50']),
    ('Arranhador Braco', 'Arranhador Adesivo', 'Cor', ['Bege', 'Cinza', 'Preto'], None),
    ('Arranhador Parede', 'Arranhador Adesivo', 'Modelo', ['Reto', 'Curvo'], None),
    ('Casinha Dobravel', 'Casinha / Caminha', 'Cor', ['Cinza', 'Azul'], None),
    ('Caminha Redonda', 'Casinha / Caminha', 'Tamanho', ['P', 'M', 'G', 'GG'], None),
    ('Comedouro Automatico', 'Comedouro Automatico', 'Modelo', ['Simples', 'Wi-Fi'], None),
    ('Fonte Bebedouro', 'Bebedouro', 'Cor', ['Branco', 'Azul'], None),
    ('Peitoral Ajustavel', 'Coleira / Peitoral', 'Tamanho', ['P', 'M', 'G', 'XG', 'XGG'], None),
]

CANAIS = [('MERCADO LIVRE', 'ML', 1.0), ('Shopee', 'SHOPEE', 0.7), ('Magalu', 'MAGALU', 0.15),
          ('Site', 'SITE', 0.12), ('AMAZON', 'AMAZON', 0.2), ('SHEIN', 'SHEIN', 0.08)]

CUSTOS = [
    ('FIXO GERAL', 'Aluguel', 'Imobiliaria Exemplo', 6500, 6500),
    ('FIXO GERAL', 'Energia', 'Companhia de Luz Ficticia', 700, 1100),
    ('FIXO GERAL', 'Internet', 'Provedor Teste', 250, 250),
    ('FIXO COM PESSOAL', 'Salarios', 'Folha (ficticia)', 18000, 21000),
    ('FIXO COM PESSOAL', 'FGTS', 'Caixa (ficticio)', 1400, 1700),
    ('FIXO COM PESSOAL', 'Vale transporte', 'Operadora VT Teste', 900, 1200),
    ('ADS', 'Marketing', 'Plataforma de Anuncios A', 3000, 9000),
    ('ADS ML RICAPET', 'Marketing', 'Mercado Ads (ficticio)', 4000, 12000),
    ('MARKETING', 'Marketing', 'Agencia Exemplo', 1500, 2500),
    ('FRETE', 'Frete', 'Transportadora Teste', 800, 3000),
    ('IMPOSTOS', 'Simples Nacional', 'Receita (ficticio)', 6000, 14000),
    ('IOF', 'IOF', 'Banco Ficticio', 30, 120),
    ('SISTEMA', 'Sistema', 'ERP Exemplo', 450, 450),
    ('SISTEMA', 'Sistema', 'Hub de Integracao Teste', 300, 300),
    ('PRONAMP', 'Emprestimo', 'Banco Ficticio', 2500, 2500),
    ('TARIFA BANCARIA', 'Tarifa', 'Banco Ficticio', 80, 200),
    ('RETIRADAS/APORTES', 'Retirada', 'Socio A (ficticio)', 5000, 8000),
]


def fmt_aba(ano, mes, i):
    """Alterna os formatos de nome de aba vistos no arquivo real."""
    yy = ano % 100
    nome = MESES_NOME[mes - 1]
    formatos = [f'{yy} - {nome}', f'{yy}- {nome}', f'{yy}-{nome}', f'{yy}- {nome} ']
    return formatos[0] if ano >= 2026 or i % 4 == 0 else formatos[i % 4]


def gerar_mes(ws, ano, mes, rnd, com_erro, prejuizo, retirada_texto):
    sazon = 1.0 + 0.25 * (mes in (11, 12)) - 0.1 * (mes in (2, 3)) + rnd.uniform(-0.08, 0.08)

    # ----- tabela de produtos (B..K)
    r = 2
    fat_cat = {}
    luc_cat = {}
    tot = {'u': 0, 'e': 0, 'c': 0, 'l': 0}
    for nome_b, nome_m, peso in CANAIS:
        ws.cell(r, 2, nome_b)
        r += 1
        cn = {'u': 0, 'e': 0, 'c': 0, 'l': 0}
        for prod, cat, tipo, variacoes, subgrupos in PRODUTOS:
            if nome_m in ('SHEIN', 'MAGALU') and rnd.random() < 0.4:
                continue
            ws.cell(r, 2, prod)
            ws.cell(r, 3, tipo)
            for j, h in enumerate(['Unidades', 'Valor Unitario', 'Valor de entrada', 'Custo Unitario',
                                   'Custo Total', 'Lucro por unidade', 'Lucro total']):
                ws.cell(r, 4 + j, h)
            r += 1
            pu = {'u': 0, 'e': 0, 'c': 0, 'l': 0}
            preco_base = rnd.uniform(35, 180)
            for sg in (subgrupos or [None]):
                fator_sg = 1 if sg is None else (1 + 0.35 * (subgrupos.index(sg)))
                for k, var in enumerate(variacoes):
                    unid = max(0, int(rnd.gauss(60, 30) * peso * sazon / (len(subgrupos or [1]))))
                    vu = round(preco_base * fator_sg * rnd.uniform(0.95, 1.05), 2)
                    cu = round(vu * rnd.uniform(0.38, 0.62), 2)
                    if nome_m == 'SHEIN' and k == 0 and prejuizo:
                        cu = round(vu * 1.1, 2)  # uma variacao no prejuizo
                    ve, ct = unid * vu, unid * cu
                    lt = ve - ct
                    ws.cell(r, 2, sg if (sg and k == 0) else None)
                    ws.cell(r, 3, var)
                    for j, val in enumerate([unid, vu, ve, cu, ct, round(vu - cu, 2), lt, (lt / ve if ve else 0)]):
                        ws.cell(r, 4 + j, val)
                    r += 1
                    for a, b in (('u', unid), ('e', ve), ('c', ct), ('l', lt)):
                        pu[a] += b
            ws.cell(r, 4, pu['u']); ws.cell(r, 6, pu['e']); ws.cell(r, 8, pu['c']); ws.cell(r, 10, pu['l'])
            ws.cell(r, 11, pu['l'] / pu['e'] if pu['e'] else 0)
            r += 2
            fat_cat.setdefault(nome_m, {}).setdefault(cat, 0)
            fat_cat[nome_m][cat] += pu['e']
            luc_cat.setdefault(nome_m, {}).setdefault(cat, 0)
            luc_cat[nome_m][cat] += pu['l']
            for a in cn:
                cn[a] += pu[a]
        ws.cell(r, 2, 'TOTAL')
        ws.cell(r, 4, cn['u']); ws.cell(r, 6, cn['e']); ws.cell(r, 8, cn['c']); ws.cell(r, 10, cn['l'])
        ws.cell(r, 11, cn['l'] / cn['e'] if cn['e'] else 0)
        r += 3
        for a in tot:
            tot[a] += cn[a]
    ws.cell(r, 2, 'TOTAL GERAL')
    ws.cell(r, 4, tot['u']); ws.cell(r, 6, tot['e']); ws.cell(r, 8, tot['c']); ws.cell(r, 10, tot['l'])
    ws.cell(r, 11, tot['l'] / tot['e'])

    # ----- canal x categoria (M/N/O faturamento, P/Q/R lucro)
    for col, fonte in ((13, fat_cat), (16, luc_cat)):
        rr = 3
        total = 0
        for _, nome_m, _ in CANAIS:
            cats = fonte.get(nome_m, {})
            primeiro = True
            for cat, val in cats.items():
                if primeiro:
                    ws.cell(rr, col, nome_m)
                    primeiro = False
                ws.cell(rr, col + 1, cat)
                ws.cell(rr, col + 2, val)
                rr += 1
            ws.cell(rr, col + 2, sum(cats.values()))  # subtotal: N vazio
            total += sum(cats.values())
            rr += 1
            if nome_m == 'SHOPEE':
                ws.cell(rr, col + 1, 'ok')  # lixo que aparece na planilha real
                ws.cell(rr, col + 2, '#VALUE!')
                rr += 1
        ws.cell(rr, col, 'Total')
        ws.cell(rr, col + 2, total)

    # ----- lancamentos de custos (Z..AE, cabecalho na linha 5)
    for j, h in enumerate(['Data', 'Valor', 'REFERENCIA', 'CATEGORIA', 'ONDE', 'OBS']):
        ws.cell(5, 26 + j, h)
    rr = 6
    custos = 0
    for ref, cat, onde, vmin, vmax in CUSTOS:
        for _ in range(rnd.randint(1, 4) if ref in ('ADS', 'FRETE', 'IOF', 'TARIFA BANCARIA') else 1):
            v = round(rnd.uniform(vmin, vmax) * (5.0 if prejuizo and ref in ('ADS', 'ADS ML RICAPET', 'FRETE') else 1), 2)
            ws.cell(rr, 26, dt.datetime(ano, mes, rnd.randint(1, 28)))
            ws.cell(rr, 27, v)
            ws.cell(rr, 28, ref)
            ws.cell(rr, 29, cat)
            ws.cell(rr, 30, onde if rr != 9 else 'Contabilidade Beta 12.345.678/0001-90')
            ws.cell(rr, 31, 'CPF 123.456.789-00 (ficticio)' if rr % 5 == 0 else None)
            custos += v
            rr += 1

    # ----- investimentos (AG/AH)
    ws.cell(5, 33, 'Investimentos')
    if mes % 3 == 0:
        ws.cell(6, 33, round(rnd.uniform(800, 4000), 2)); ws.cell(6, 34, 'Moveis do escritorio (exemplo)')
        ws.cell(7, 33, round(rnd.uniform(300, 1500), 2)); ws.cell(7, 34, 'Impressora de etiquetas (exemplo)')

    # ----- resumo (T/U)
    fat_bruto = tot['e'] / 0.82
    lucro_bruto = tot['l']
    lucro_liq = lucro_bruto - custos
    bonus = round(max(0, lucro_liq) * 0.05, 2)
    linhas = [('Faturamento Bruto', fat_bruto), ('Lucro Bruto', lucro_bruto), ('Custos Mensais', custos),
              ('Lucro Liquido', lucro_liq),
              ('Retirada socios ', 'JÁ CONSIDERADO NO CUSTO' if retirada_texto else 6000.0),
              ('Bonus Equipe Exemplo', bonus), ('Lucro empresa', lucro_liq - bonus)]
    for i, (rot, val) in enumerate(linhas):
        ws.cell(3 + i, 20, rot)
        ws.cell(3 + i, 21, val)

    # ----- quadro de metas (T12:X18)
    for j, h in enumerate(['Valor', 'porcentagem', 'Meta', 'Resultado']):
        ws.cell(12, 21 + j, h)
    taxas = fat_bruto - tot['e']
    cmv = tot['c']
    contrib = fat_bruto - taxas - cmv
    metas = [('Faturamento Bruto', fat_bruto, 1.0, 1.0, False), ('Taxas Mkt Place', taxas, taxas / fat_bruto, 0.20, True),
             ('CMV', cmv, cmv / fat_bruto, 0.36, True), ('Contribuição', contrib, contrib / fat_bruto, 0.42, False),
             ('Custos fixos', custos, custos / fat_bruto, 0.25, True), ('Lucro liquido', lucro_liq, lucro_liq / fat_bruto, 0.12, False)]
    for i, (rot, val, pct, meta, custo) in enumerate(metas):
        ws.cell(13 + i, 20, rot)
        ws.cell(13 + i, 21, val)
        ws.cell(13 + i, 22, pct)
        ws.cell(13 + i, 23, meta)
        ws.cell(13 + i, 24, (meta - pct) if custo else (pct - meta))

    # ----- conferencias (T45, T58)
    ws.cell(45, 19, 'Conferencia faturamento')
    ws.cell(45, 20, 'Algo Errado' if com_erro else 'OK')
    ws.cell(58, 19, 'Conferencia custos')
    ws.cell(58, 20, 'OK')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--sem-valores', action='store_true')
    args = ap.parse_args()
    rnd = random.Random(42)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet('Planilha1').cell(1, 1, 'aba solta (ignorar)')
    meses = []
    a, m = 2025, 7
    for _ in range(15):
        meses.append((a, m))
        m += 1
        if m > 12:
            a, m = a + 1, 1
    for i, (ano, mes) in enumerate(meses):
        ws = wb.create_sheet(fmt_aba(ano, mes, i))
        gerar_mes(ws, ano, mes, rnd, com_erro=(ano, mes) == (2026, 9), prejuizo=(ano, mes) == (2026, 2),
                  retirada_texto=(i % 2 == 0))
    q = wb.create_sheet('Quantidad mes a mes')
    q.cell(1, 1, 'Cor')
    if args.sem_valores:
        # Simula planilha salva por script: tudo vira formula sem valor em cache.
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, (int, float)):
                        c.value = f'={c.value}+0'
        destino = os.path.join(PASTA, 'Fechamento_Exemplo_SemValores.xlsx')
    else:
        destino = os.path.join(PASTA, 'Fechamento_Exemplo.xlsx')
    wb.save(destino)
    print('Gerado:', destino)


if __name__ == '__main__':
    main()
